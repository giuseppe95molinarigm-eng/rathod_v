"""Build the KDP-ready 6 x 9 in Word manuscript from manuscript/*.md.

Usage: python3 build_book.py OUT.docx [pages.json]

Source markup (one file per front-matter piece, part or chapter):
  @title / @copyright / @contents       special front-matter pages
  @front Title                          unnumbered front-matter section
  @part One | Title | intro text        part opener page
  @chapter 1 | Title                    chapter opener
  > text                                italic chapter intro / epigraph
  ## / ###                              section headings
  - item / [ ] item / 1. item           bullets, checklist, numbered list
  | a | b |                             table (first row is the header)
  :::kind Title ... :::                 box (kind: summary, try, evidence, note)
  **bold**, *italic*                    inline formatting
pages.json maps TOC entry ids to printed page numbers (filled by make_book.sh).
"""

import glob
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.environ.get("BOOK_SRC", os.path.join(ROOT, "manuscript"))
# PAPER=letter or PAPER=a4 lays out a standalone printable (the Printables
# PDF): that paper size, even margins, no running heads, no blank pages
PAPER = {"letter": (8.5, 11), "a4": (8.27, 11.69)}.get(os.environ.get("PAPER", ""))
OUT = sys.argv[1]
PAGES = json.load(open(sys.argv[2])) if len(sys.argv) > 2 and os.path.exists(sys.argv[2]) else {}

TITLE = "The Menopause Diet for Busy Women"
SUBTITLE = ("Quick, Delicious Recipes and a 12-Week Meal Plan to Boost Energy, "
            "Sleep Better and Feel Like Yourself Again")
AUTHOR = "Neha Menon"

BODY_FONT = "Cambria"
HEAD_FONT = "Calibri"
ACCENT_HEX = "5A3E5D"
ACCENT = RGBColor.from_string(ACCENT_HEX)
GREY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x26, 0x26, 0x26)
BOX_FILL = {"summary": "F3EFF3", "try": "F6F1EA", "evidence": "EEF2F0", "note": "F2F2F2",
            "why": "EEF2F0", "swaps": "F2F2F2", "important": "F7ECEC"}
BOX_LABEL = {"summary": "Chapter Summary", "try": "Try This Week",
             "evidence": "What the Evidence Says", "note": "Good to Know",
             "why": "Why It Helps", "swaps": "Swaps & Variations", "important": "Important"}

doc = Document()
styles = doc.styles


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

def set_font(style, name, size, bold=False, italic=False, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def para_fmt(style, before=0, after=6, line=1.2, keep_next=False, align=None,
             indent=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.left_indent = indent


def new_style(name, base="Normal"):
    s = styles.add_style(name, 1)
    s.base_style = styles[base]
    return s


def outline_level(style, level):
    ppr = style.element.get_or_add_pPr()
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    ppr.append(el)


set_font(styles["Normal"], BODY_FONT, 10.5)
para_fmt(styles["Normal"], after=6, line=1.22, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

set_font(styles["Heading 1"], HEAD_FONT, 22, bold=True, color=ACCENT)
para_fmt(styles["Heading 1"], before=0, after=6, line=1.0, keep_next=True,
         align=WD_ALIGN_PARAGRAPH.LEFT)
set_font(styles["Heading 2"], HEAD_FONT, 13, bold=True, color=ACCENT)
para_fmt(styles["Heading 2"], before=14, after=4, line=1.1, keep_next=True,
         align=WD_ALIGN_PARAGRAPH.LEFT)
set_font(styles["Heading 3"], HEAD_FONT, 11, bold=True, color=DARK)
para_fmt(styles["Heading 3"], before=8, after=2, line=1.1, keep_next=True,
         align=WD_ALIGN_PARAGRAPH.LEFT)

for name in ("List Bullet", "List Bullet 2", "List Number"):
    set_font(styles[name], BODY_FONT, 10.5)
    para_fmt(styles[name], after=3, line=1.18, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Chapter Label")
set_font(s, HEAD_FONT, 10, bold=True, color=GREY)
para_fmt(s, before=54, after=4, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Part Label")
set_font(s, HEAD_FONT, 12, bold=True, color=GREY)
para_fmt(s, before=150, after=8, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Part Title")
set_font(s, HEAD_FONT, 26, bold=True, color=ACCENT)
para_fmt(s, after=18, line=1.0, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
outline_level(s, 0)

s = new_style("Part Intro")
set_font(s, BODY_FONT, 11, italic=True, color=GREY)
para_fmt(s, after=6, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Front Title", "Heading 1")
para_fmt(s, before=54, after=14, line=1.0, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Intro")
set_font(s, BODY_FONT, 11.5, italic=True, color=GREY)
para_fmt(s, before=4, after=14, line=1.3, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Scene")
set_font(s, BODY_FONT, 10.5, italic=True, color=DARK)
para_fmt(s, after=6, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Inches(0.25))
s.paragraph_format.right_indent = Inches(0.25)

s = new_style("Bridge")
set_font(s, BODY_FONT, 11, italic=True, color=ACCENT)
para_fmt(s, before=18, after=0, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Box Title")
set_font(s, HEAD_FONT, 10.5, bold=True, color=ACCENT)
para_fmt(s, after=3, line=1.1, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Box Text")
set_font(s, BODY_FONT, 10)
para_fmt(s, after=3, line=1.18, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Check Item")
set_font(s, BODY_FONT, 10.5)
para_fmt(s, after=3, line=1.18, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Inches(0.28))
s.paragraph_format.first_line_indent = Inches(-0.22)

s = new_style("Recipe Title")
set_font(s, HEAD_FONT, 18, bold=True, color=ACCENT)
para_fmt(s, before=6, after=3, line=1.05, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Recipe Tags")
set_font(s, HEAD_FONT, 8.5, bold=True, color=GREY)
para_fmt(s, after=8, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Ingredient")
set_font(s, BODY_FONT, 10)
para_fmt(s, after=1.5, line=1.12, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Inches(0.18))
s.paragraph_format.first_line_indent = Inches(-0.14)

s = new_style("Ingredient Group")
set_font(s, HEAD_FONT, 9, bold=True, color=DARK)
para_fmt(s, before=4, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Table Text")
set_font(s, BODY_FONT, 9)
para_fmt(s, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)

s = new_style("Center Small")
set_font(s, BODY_FONT, 9, color=GREY)
para_fmt(s, after=4, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Copyright")
set_font(s, BODY_FONT, 8.5, color=DARK)
para_fmt(s, after=6, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT)

# title page colors follow the cover: deep green title, coral second line
COVER_GREEN = RGBColor.from_string("042F1C")
COVER_CORAL = RGBColor.from_string("E24F2E")

s = new_style("Title Big")
set_font(s, HEAD_FONT, 32, bold=True, color=COVER_GREEN)
para_fmt(s, before=110, after=18, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Title Sub")
set_font(s, HEAD_FONT, 13, bold=True, color=COVER_GREEN)
para_fmt(s, after=40, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)

s = new_style("Title Author")
set_font(s, HEAD_FONT, 13, bold=True, color=DARK)
para_fmt(s, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)

if PAPER:
    # printables: tighter type so each meal-plan week fits on one page
    for name, size, line in (("Normal", 9.5, 1.15), ("Table Text", 8.5, 1.05),
                             ("Check Item", 9.5, 1.12), ("Box Text", 9, 1.1)):
        set_font(styles[name], BODY_FONT, size)
        styles[name].paragraph_format.line_spacing = line
    styles["Normal"].paragraph_format.space_after = Pt(4)

toc1 = styles["TOC 1"] if "TOC 1" in [x.name for x in styles] else new_style("TOC 1")
set_font(toc1, BODY_FONT, 10.5)
para_fmt(toc1, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
toc1.paragraph_format.tab_stops.add_tab_stop(Inches(4.6), WD_TAB_ALIGNMENT.RIGHT,
                                             WD_TAB_LEADER.DOTS)

# Settings: mirrored margins, odd/even headers, hyphenation
settings = doc.settings.element
for tag in ("w:mirrorMargins", "w:evenAndOddHeaders"):
    if not PAPER:
        settings.insert(0, OxmlElement(tag))
hyph = OxmlElement("w:autoHyphenation")
settings.append(hyph)
lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "en-US")
styles["Normal"].element.get_or_add_rPr().append(lang)

# ---------------------------------------------------------------------------
# Section / header helpers
# ---------------------------------------------------------------------------


def setup_page(sec):
    if PAPER:
        sec.page_width, sec.page_height = Inches(PAPER[0]), Inches(PAPER[1])
        for side in ("top", "bottom", "left", "right"):
            setattr(sec, side + "_margin", Inches(0.6))
        sec.header_distance = sec.footer_distance = Inches(0.4)
        return
    sec.page_width = Inches(6)
    sec.page_height = Inches(9)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.8)    # inside / gutter
    sec.right_margin = Inches(0.6)   # outside
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)


def page_num_format(sec, fmt=None, start=None):
    sectpr = sec._sectPr
    el = sectpr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType")
        sectpr.append(el)
    if fmt:
        el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))


def add_field(paragraph, code, cached="1", size=9, color=GREY, font=HEAD_FONT):
    def mk_run():
        r = paragraph.add_run()
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
        return r
    r = mk_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r._r.append(fc)
    r = mk_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {code} "
    r._r.append(it)
    r = mk_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r._r.append(fc)
    r = mk_run()
    r.text = cached
    r = mk_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r._r.append(fc)


def clear(part):
    part.is_linked_to_previous = False
    for p in part.paragraphs:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    return part.paragraphs[0]


def header_text(part, text, align):
    p = clear(part)
    p.alignment = align
    if text:
        r = p.add_run(text.upper())
        r.font.name = HEAD_FONT
        r.font.size = Pt(7.5)
        r.font.color.rgb = GREY
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), "30")
        r._r.get_or_add_rPr().append(sp)


def footer_number(part, show=True):
    p = clear(part)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if show:
        add_field(p, "PAGE")


def configure_section(sec, recto_text, verso_text, numbers=True, first_blank=True):
    """Headers: odd (recto) = recto_text, even (verso) = verso_text.
    The first page of the section has no running header."""
    setup_page(sec)
    if PAPER:
        recto_text = verso_text = ""
    sec.different_first_page_header_footer = first_blank
    header_text(sec.header, recto_text, WD_ALIGN_PARAGRAPH.RIGHT)
    header_text(sec.even_page_header, verso_text, WD_ALIGN_PARAGRAPH.LEFT)
    header_text(sec.first_page_header, "", WD_ALIGN_PARAGRAPH.CENTER)
    footer_number(sec.footer, numbers)
    footer_number(sec.even_page_footer, numbers)
    footer_number(sec.first_page_footer, numbers)


first_section_used = False


def new_section(start=WD_SECTION.ODD_PAGE):
    global first_section_used
    if not first_section_used:
        first_section_used = True
        return doc.sections[0]
    sec = doc.add_section(WD_SECTION.NEW_PAGE if PAPER else start)
    # a new section copies the previous one's page numbering; keep the format
    # but never the restart value, so numbering runs on continuously
    el = sec._sectPr.find(qn("w:pgNumType"))
    if el is not None and el.get(qn("w:start")) is not None:
        del el.attrib[qn("w:start")]
    return sec


# ---------------------------------------------------------------------------
# Inline + block helpers
# ---------------------------------------------------------------------------

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
MARK = re.compile(r"(==.+?==)")
# ==text== marks author's personal moments for the client to confirm;
# set HIGHLIGHT=0 to build without the yellow highlight
SHOW_MARKS = os.environ.get("HIGHLIGHT", "1") != "0"


REFS = []          # (key, kind, text) page references used by the indexes
REF = re.compile(r"\{\{([rp]):([^}]+)\}\}")


def resolve_refs(text):
    def sub(m):
        kind, target = m.group(1), m.group(2).strip()
        key = f"ref|{kind}|{target}"
        if key not in [k for k, _, _ in REFS]:
            REFS.append((key, kind, target))
        return str(PAGES.get(key, "00"))
    out = REF.sub(sub, text)
    # collapse repeated page numbers in a list ("13, 13, 41" -> "13, 41")
    return re.sub(r"\b(\d+)(, \1\b)+", r"\1", out)


def add_inline(p, text, size=None, color=None, bold=None):
    text = resolve_refs(text).replace("[ ] ", "\u2610 ")
    for part in MARK.split(text):
        if part.startswith("==") and part.endswith("==") and len(part) > 4:
            start = len(p.runs)
            _add_runs(p, part[2:-2], size, color, bold)
            if SHOW_MARKS:
                for r in p.runs[start:]:
                    hl = OxmlElement("w:highlight")
                    hl.set(qn("w:val"), "yellow")
                    r._r.get_or_add_rPr().append(hl)
        elif part:
            _add_runs(p, part, size, color, bold)
    return p


def _add_runs(p, text, size=None, color=None, bold=None):
    if "\u2610" in text:
        # checkbox glyphs get a font that has them (avoids color-emoji fallback)
        pieces = text.split("\u2610")
        for k, piece in enumerate(pieces):
            if k:
                r = p.add_run("\u2610")
                r.font.name = "DejaVu Sans"
            if piece:
                _add_runs(p, piece, size, color, bold)
        return p
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**"):
            r = p.add_run(chunk[2:-2])
            r.bold = True
        elif chunk.startswith("*"):
            r = p.add_run(chunk[1:-1])
            r.italic = True
        else:
            r = p.add_run(chunk)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        if bold:
            r.bold = True
    return p


def shade_cell(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_borders(cell, left=None):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side == "left" and left:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "24")
            el.set(qn("w:color"), left)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcpr.append(b)


def table_setup(t, widths, borders=True, margins=(70, 110)):
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", margins[0]), ("left", margins[1]),
                      ("bottom", margins[0]), ("right", margins[1])):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if borders and side in ("top", "bottom", "insideH"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "BFB3C0")
        else:
            el.set(qn("w:val"), "nil")
        bd.append(el)
    tblpr.append(bd)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Inches(widths[i])


def row_flags(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if header:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)


TEXT_WIDTH = PAPER[0] - 1.2 if PAPER else 6 - 0.8 - 0.6


def add_table(rows):
    header, body = rows[0], rows[1:]
    n = len(header)
    # optional fixed width for a column: "Day@0.95" in the header (inches)
    fixed = {}
    for k, h in enumerate(header):
        m = re.match(r"(.*)@([\d.]+)$", h)
        if m:
            header[k] = m.group(1)
            fixed[k] = float(m.group(2))
    # width proportional to the longest text in each column
    lens = [max(len(r[i]) if i < len(r) else 0 for r in rows) for i in range(n)]
    weights = [max(8, min(l, 70)) ** 0.75 for l in lens]
    widths = [TEXT_WIDTH * w / sum(weights) for w in weights]
    # make sure no column is squeezed below a readable minimum
    minimum = 0.85 if n <= 4 else 0.7
    for _ in range(3):
        short = [i for i, w in enumerate(widths) if w < minimum]
        if not short:
            break
        extra = sum(minimum - widths[i] for i in short)
        wide = [i for i in range(n) if i not in short]
        total_wide = sum(widths[i] for i in wide)
        for i in short:
            widths[i] = minimum
        for i in wide:
            widths[i] -= extra * widths[i] / total_wide
    if fixed:
        free = [k for k in range(n) if k not in fixed]
        rest = TEXT_WIDTH - sum(fixed.values())
        tot = sum(widths[k] for k in free) or 1
        widths = [fixed[k] if k in fixed else rest * widths[k] / tot for k in range(n)]
    t = doc.add_table(rows=1 + len(body), cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_setup(t, widths)
    for ri, row in enumerate([header] + body):
        row_flags(t.rows[ri], header=(ri == 0))
        if ri and all(not c.strip() for c in row[1:]) and n > 2:
            # blank rows are for writing in: give them room
            trpr = t.rows[ri]._tr.get_or_add_trPr()
            h = OxmlElement("w:trHeight")
            h.set(qn("w:val"), str(int((0.34 if n >= 6 else 0.78) * 1440)))
            h.set(qn("w:hRule"), "atLeast")
            trpr.append(h)
        for ci in range(n):
            c = t.rows[ri].cells[ci]
            p = c.paragraphs[0]
            p.style = styles["Table Text"]
            txt = row[ci] if ci < len(row) else ""
            if ri == 0:
                shade_cell(c, ACCENT_HEX)
                r = p.add_run(txt)
                r.bold = True
                r.font.name = HEAD_FONT
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            else:
                if ri % 2 == 0:
                    shade_cell(c, "F7F4F7")
                parts = txt.split(" <br> ")
                for k, part in enumerate(parts):
                    if k:
                        p.add_run().add_break()
                    add_inline(p, part)
                if ci == 0:
                    for r in p.runs:
                        r.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_box(kind, title, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_setup(t, [TEXT_WIDTH], borders=False, margins=(110, 160))
    row_flags(t.rows[0])
    c = t.rows[0].cells[0]
    shade_cell(c, BOX_FILL.get(kind, "F2F2F2"))
    cell_borders(c, left=ACCENT_HEX)
    p = c.paragraphs[0]
    p.style = styles["Box Title"]
    label = BOX_LABEL.get(kind, "")
    # evidence and note boxes always carry their family label, so readers
    # can spot the box types promised in "How to Use This Book"
    if title and kind in ("evidence", "note", "important") and title.lower() != label.lower():
        text = f"{label}: {title}"
    else:
        text = title or label
    p.add_run(text.upper())
    for line in lines:
        if line.startswith("- "):
            q = c.add_paragraph(style="Box Text")
            q.paragraph_format.left_indent = Inches(0.18)
            q.paragraph_format.first_line_indent = Inches(-0.14)
            add_inline(q, "•  " + line[2:])
        elif re.match(r"\d+\. ", line):
            q = c.add_paragraph(style="Box Text")
            q.paragraph_format.left_indent = Inches(0.22)
            q.paragraph_format.first_line_indent = Inches(-0.18)
            num, _, rest = line.partition(" ")
            add_inline(q, num + "  " + rest)
        elif line.startswith("[ ] "):
            q = c.add_paragraph(style="Box Text")
            q.paragraph_format.left_indent = Inches(0.22)
            q.paragraph_format.first_line_indent = Inches(-0.2)
            r = q.add_run("☐  ")
            r.font.name = "DejaVu Sans"
            add_inline(q, line[4:])
        else:
            add_inline(c.add_paragraph(style="Box Text"), line)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


_bm = [100]


def bookmark(paragraph, name):
    _bm[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bm[0]))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bm[0]))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


TOC_ENTRIES = []   # (id, text, is_part)


def register(paragraph, text, is_part=False):
    tid = f"_Toc{len(TOC_ENTRIES) + 100000}"
    bookmark(paragraph, tid)
    TOC_ENTRIES.append((tid, text, is_part))


# ---------------------------------------------------------------------------
# Special pages
# ---------------------------------------------------------------------------

def title_pages():
    sec = new_section()
    setup_page(sec)
    sec.different_first_page_header_footer = True
    for part in (sec.header, sec.even_page_header, sec.first_page_header):
        header_text(part, "", WD_ALIGN_PARAGRAPH.CENTER)
    for part in (sec.footer, sec.even_page_footer, sec.first_page_footer):
        footer_number(part, False)
    p = doc.add_paragraph(style="Title Big")
    first, _, second = TITLE.partition(" for ")
    p.add_run(first)
    p.add_run().add_break()
    r = p.add_run("for " + second)
    r.font.color.rgb = COVER_CORAL
    doc.add_paragraph(SUBTITLE, style="Title Sub")
    doc.add_paragraph(AUTHOR, style="Title Author")


def copyright_page(lines):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    for _ in range(8):
        doc.add_paragraph(style="Copyright")
    for line in lines:
        add_inline(doc.add_paragraph(style="Copyright"), line)


def contents_page():
    sec = new_section()
    configure_section(sec, "Contents", TITLE)
    page_num_format(sec, fmt="lowerRoman", start=None)
    h = doc.add_paragraph("Contents", style="Front Title")
    # TOC field with cached, clickable entries
    p = doc.add_paragraph(style="TOC 1")
    r = p.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r._r.append(fc)
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-1" \\h \\z \\u '
    r._r.append(it)
    r = p.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r._r.append(fc)
    return p


def fill_toc(first_p):
    p = first_p
    for i, (tid, text, is_part) in enumerate(TOC_ENTRIES):
        if i > 0:
            new = OxmlElement("w:p")
            p._p.addnext(new)
            from docx.text.paragraph import Paragraph
            p = Paragraph(new, first_p._parent)
            p.style = styles["TOC 1"]
        if is_part:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("w:anchor"), tid)
        hl.set(qn("w:history"), "1")
        p._p.append(hl)

        def run(txt, bold=False):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            if bold:
                rpr.append(OxmlElement("w:b"))
            r.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = txt
            r.append(t)
            return r

        hl.append(run(text, bold=is_part))
        tab = OxmlElement("w:r")
        tab.append(OxmlElement("w:tab"))
        hl.append(tab)
        # PAGEREF field
        for kind, payload in (("begin", None), ("instr", f" PAGEREF {tid} \\h "),
                              ("separate", None), ("text", str(PAGES.get(tid, "0"))),
                              ("end", None)):
            r = OxmlElement("w:r")
            if kind == "instr":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = payload
                r.append(it)
            elif kind == "text":
                r = run(payload, bold=is_part)
            else:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r.append(fc)
            hl.append(r)
    # close the TOC field
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r.append(fc)
    p._p.append(r)


CURRENT = {"chapter": ""}


def recipe_title(text):
    # each recipe opens on a left-hand page so its two pages face each other.
    # LibreOffice drops mirrored margins in even-page sections, so recipes use
    # plain new-page sections and make_book.sh lists the ones that need a
    # blank page in front (PAGES["_pad"]) to land on the left.
    if text in PAGES.get("_pad", []):
        sec = new_section(WD_SECTION.NEW_PAGE)
        configure_section(sec, "", "", numbers=False, first_blank=False)
        doc.add_paragraph()
    sec = new_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, CURRENT["chapter"], TITLE, first_blank=False)
    doc.add_paragraph(text, style="Recipe Title")
    RECIPES.append(text)
    SPREADS.append([text, CURRENT["chapter"]])


SNACK = {"start": None}
NEWPAGE = {"on": False}


def close_snack():
    """Keep a finished snack's paragraphs together on one page."""
    if SNACK["start"] is None:
        return
    paras = doc.paragraphs[SNACK["start"]:]
    for q in paras[:-1]:
        q.paragraph_format.keep_with_next = True
    SNACK["start"] = None


def snack_title(text):
    close_snack()
    SNACK["start"] = len(doc.paragraphs)
    p = doc.add_paragraph(text, style="Recipe Title")
    p.paragraph_format.space_before = Pt(20)
    for r in p.runs:
        r.font.size = Pt(14)
    RECIPES.append(text)


def snack_info(text):
    p = doc.add_paragraph(style="Recipe Tags")
    r = p.add_run("  \u00b7  ".join(t.strip().upper() for t in text.split("|")))
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "20")
    r._r.get_or_add_rPr().append(sp)


def recipe_tags(text):
    p = doc.add_paragraph(style="Recipe Tags")
    r = p.add_run("  \u00b7  ".join(t.strip().upper() for t in text.split("|")))
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "20")
    r._r.get_or_add_rPr().append(sp)


def recipe_meta(text):
    items = [x.strip() for x in text.split("|")]
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_setup(t, [TEXT_WIDTH / len(items)] * len(items), borders=False,
                margins=(70, 80))
    row_flags(t.rows[0])
    for c, item in zip(t.rows[0].cells, items):
        shade_cell(c, BOX_FILL["summary"])
        label, _, value = item.partition(" ")
        p = c.paragraphs[0]
        p.style = styles["Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label.upper() + "\n")
        r.bold = True
        r.font.name = HEAD_FONT
        r.font.size = Pt(7.5)
        r.font.color.rgb = ACCENT
        r = p.add_run(value)
        r.font.size = Pt(9.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


RECIPES = []
SPREADS = []           # (recipe, chapter) for recipes that open a spread


# ---------------------------------------------------------------------------
# Markup parser
# ---------------------------------------------------------------------------

def render_body(lines):
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            flush()
            i += 1
            continue
        if line.startswith(":::"):
            flush()
            m = re.match(r":::(\w+)\s*(.*)", line)
            kind, title = m.group(1), m.group(2).strip()
            box = []
            i += 1
            while i < len(lines) and not lines[i].startswith(":::"):
                if lines[i].strip():
                    box.append(lines[i].rstrip())
                i += 1
            add_box(kind, title, box)
            i += 1
            continue
        if line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(rows)
            continue
        if line.strip() == "@newpage":
            flush()
            NEWPAGE["on"] = True
            i += 1
            continue
        if line.strip() == "@qr":
            # QR code for the printables link (print only)
            flush()
            q = doc.add_paragraph(style="Center Small")
            q.paragraph_format.keep_with_next = False
            q.add_run().add_picture(os.path.join(ROOT, "assets", "qr_printables.png"),
                                    width=Inches(1.1))
            i += 1
            continue
        if line.startswith("@recipe "):
            flush()
            close_snack()
            recipe_title(line[8:].strip())
        elif line.startswith("@snack "):
            flush()
            snack_title(line[7:].strip())
        elif line.startswith("@info "):
            flush()
            snack_info(line[6:])
        elif line.startswith("@tags "):
            flush()
            recipe_tags(line[6:])
        elif line.startswith("@meta "):
            flush()
            recipe_meta(line[6:])
        elif line.startswith("* "):
            flush()
            q = doc.add_paragraph(style="Ingredient")
            r = q.add_run("\u25aa  ")
            r.font.color.rgb = ACCENT
            r.font.name = "DejaVu Sans"
            add_inline(q, line[2:])
        elif line.startswith("#### "):
            flush()
            doc.add_paragraph(line[5:], style="Ingredient Group")
        elif line.startswith("### "):
            flush()
            doc.add_paragraph(line[4:], style="Heading 3")
        elif line.startswith("## "):
            flush()
            close_snack()
            h2 = doc.add_paragraph(line[3:], style="Heading 2")
            if NEWPAGE["on"]:
                h2.paragraph_format.page_break_before = True
                NEWPAGE["on"] = False
        elif line.startswith(">> "):
            flush()
            orn = doc.add_paragraph("\u2022  \u2022  \u2022", style="Bridge")
            orn.paragraph_format.keep_with_next = True
            add_inline(doc.add_paragraph(style="Bridge"), line[3:]).paragraph_format.space_before = Pt(4)
        elif line.startswith("%% "):
            flush()
            add_inline(doc.add_paragraph(style="Scene"), line[3:])
        elif line.startswith("> "):
            flush()
            add_inline(doc.add_paragraph(style="Intro"), line[2:])
        elif line[:1] == " " and line.lstrip().startswith("- "):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet 2"), line.lstrip()[2:])
        elif line.startswith("- "):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("[ ] "):
            flush()
            p = doc.add_paragraph(style="Check Item")
            r = p.add_run("☐  ")
            r.font.name = "DejaVu Sans"
            add_inline(p, line[4:])
        elif re.match(r"\d+\. ", line):
            flush()
            num, txt = line.split(". ", 1)
            p = doc.add_paragraph(style="Check Item")
            r = p.add_run(f"{num}.  ")
            r.bold = True
            r.font.color.rgb = ACCENT
            add_inline(p, txt)
        else:
            para_buf.append(line.strip())
        i += 1
    flush()
    close_snack()


def build():
    files = sorted(glob.glob(os.path.join(SRC, "*.md")))
    toc_p = None
    body_started = False
    for path in files:
        lines = open(path, encoding="utf-8").read().splitlines()
        head = lines[0].strip()
        rest = lines[1:]
        if head == "@title":
            title_pages()
            continue
        if head == "@copyright":
            copyright_page([l for l in rest if l.strip()])
            continue
        if head == "@contents":
            toc_p = contents_page()
            continue
        if head.startswith("@front "):
            name = head[7:].strip()
            sec = new_section()
            configure_section(sec, name, TITLE)
            h = doc.add_paragraph(name, style="Front Title")
            register(h, name)
            render_body(rest)
            continue
        if head.startswith("@part "):
            num, title, *intro = [x.strip() for x in head[6:].split("|")]
            sec = new_section()
            configure_section(sec, "", "", numbers=False)
            if not body_started:
                page_num_format(sec, fmt="decimal", start=1)
                body_started = True
            doc.add_paragraph(f"PART {num.upper()}", style="Part Label")
            p = doc.add_paragraph(title, style="Part Title")
            register(p, f"Part {num}: {title}", is_part=True)
            for t in intro:
                doc.add_paragraph(t, style="Part Intro")
            render_body(rest)
            continue
        if head.startswith("@chapter "):
            num, title = [x.strip() for x in head[9:].split("|")]
            sec = new_section()
            configure_section(sec, title, TITLE)
            CURRENT["chapter"] = title
            doc.add_paragraph(f"CHAPTER {num}", style="Chapter Label")
            h = doc.add_paragraph(title, style="Heading 1")
            pbdr = OxmlElement("w:pBdr")
            b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "8")
            b.set(qn("w:space"), "6")
            b.set(qn("w:color"), ACCENT_HEX)
            pbdr.append(b)
            h._p.get_or_add_pPr().append(pbdr)
            register(h, f"{num}. {title}")
            render_body(rest)
            continue
        raise SystemExit(f"unknown header in {path}: {head}")
    if toc_p is not None:
        fill_toc(toc_p)
    doc.core_properties.title = TITLE
    doc.core_properties.author = AUTHOR
    doc.save(OUT)
    json.dump({"toc": [[t, x] for t, x, _ in TOC_ENTRIES],
               "refs": [list(r) for r in REFS], "spreads": SPREADS},
              open(os.path.splitext(OUT)[0] + ".toc.json", "w"), indent=0)
    print("saved", OUT, len(TOC_ENTRIES), "toc entries")


build()
